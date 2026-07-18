"""parse_material 工具单元测试（快速用例，不依赖 marker）。

覆盖：text / txt / md / docx / pptx / 不支持格式 / 空内容 / 分块 / 截断 / 注册表。
marker PDF 集成测试见 test_parse_material_marker.py。
"""

from __future__ import annotations

import asyncio
import tempfile

import pytest

from app.agent.context import ExecContext
from app.tools import default_registry
from app.tools.base import ToolInput
from app.tools.parse_material import CHUNK_SIZE, MAX_PARSED_CHARS, parse_material


def _ctx() -> ExecContext:
    return ExecContext(session_id="sess_test", trace_id="trace_test")


def _inp(args: dict) -> ToolInput:
    return ToolInput(tool="parse_material", session_id="sess_test", args=args)


def _run(args: dict):
    return asyncio.run(parse_material(_inp(args), _ctx()))


def _tmp(suffix: str, content: str | bytes = "") -> str:
    f = tempfile.NamedTemporaryFile(suffix=suffix, delete=False, mode="wb")
    f.write(content.encode("utf-8") if isinstance(content, str) else content)
    f.close()
    return f.name


# --- text ---


def test_text_kind():
    out = _run({"kind": "text", "source": "hello world"})
    assert out.status == "ok"
    assert out.data["text"] == "hello world"
    assert out.data["meta"]["source_type"] == "text"
    assert out.data["meta"]["chars"] == 11
    assert out.data["chunks"] == ["hello world"]


def test_empty_text_is_error():
    out = _run({"kind": "text", "source": "   "})
    assert out.status == "error"
    assert out.errors[0].code == "empty_content"


def test_unsupported_kind():
    out = _run({"kind": "video_url", "source": "x"})
    assert out.status == "error"
    assert out.errors[0].code == "unsupported_format"


# --- file: txt / md ---


def test_txt_file():
    path = _tmp(".txt", "plain text content")
    out = _run({"kind": "file", "source": path})
    assert out.status == "ok"
    assert out.data["text"] == "plain text content"
    assert out.data["meta"]["source_type"] == "txt"


def test_md_file():
    path = _tmp(".md", "# Title\n\nsome **markdown** text")
    out = _run({"kind": "file", "source": path})
    assert out.status == "ok"
    assert "markdown" in out.data["text"]
    assert out.data["meta"]["source_type"] == "md"


def test_empty_file_is_error():
    path = _tmp(".txt", "")
    out = _run({"kind": "file", "source": path})
    assert out.status == "error"
    assert out.errors[0].code == "empty_content"


def test_unsupported_extension():
    path = _tmp(".xlsx", "fake")
    out = _run({"kind": "file", "source": path})
    assert out.status == "error"
    assert out.errors[0].code == "unsupported_format"


def test_missing_file():
    out = _run({"kind": "file", "source": "/no/such/file.pdf"})
    assert out.status == "error"
    assert out.errors[0].code == "unsupported_format"


# --- file: docx / pptx ---


def test_docx_file():
    from docx import Document

    path = _tmp(".docx")
    doc = Document()
    doc.add_paragraph("docx paragraph content")
    doc.save(path)
    out = _run({"kind": "file", "source": path})
    assert out.status == "ok"
    assert "docx paragraph content" in out.data["text"]
    assert out.data["meta"]["source_type"] == "docx"


def test_pptx_file():
    from pptx import Presentation

    path = _tmp(".pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = slide.shapes.add_textbox(0, 0, 300, 100)
    tb.text_frame.text = "ppt slide content"
    prs.save(path)
    out = _run({"kind": "file", "source": path})
    assert out.status == "ok"
    assert "ppt slide content" in out.data["text"]
    assert out.data["meta"]["source_type"] == "pptx"


# --- 分块 / 截断 ---


def test_chunking_splits_long_text():
    text = "a" * (CHUNK_SIZE * 2 + 100)
    out = _run({"kind": "text", "source": text})
    assert out.status == "ok"
    assert len(out.data["chunks"]) == 3
    assert len(out.data["chunks"][0]) == CHUNK_SIZE
    assert len(out.data["chunks"][2]) == 100


def test_truncation_over_max():
    text = "b" * (MAX_PARSED_CHARS + 1000)
    out = _run({"kind": "text", "source": text})
    assert out.status == "ok"
    assert out.data["meta"]["truncated"] is True
    assert out.data["meta"]["chars"] == MAX_PARSED_CHARS


# --- 注册表 ---


def test_default_registry_has_parse_material():
    reg = default_registry()
    assert reg.has("parse_material")
    assert reg.get("parse_material").llm is False


# --- URL 抓取限量读取（防 OOM）---


class _FakeAiter:
    """模拟 httpx 响应的 aiter_bytes()：按块产出。"""

    def __init__(self, chunks: list[bytes]) -> None:
        self._chunks = chunks

    def __aiter__(self):
        self._it = iter(self._chunks)
        return self

    async def __anext__(self) -> bytes:
        try:
            return next(self._it)
        except StopIteration as exc:
            raise StopAsyncIteration from exc


def test_collect_capped_within_limit():
    from app.tools.parse_material import _collect_capped

    out, too_large = asyncio.run(_collect_capped(_FakeAiter([b"ab", b"cd"]), 10))
    assert out == b"abcd"
    assert too_large is False


def test_collect_capped_over_limit_stops_early():
    """超上限即停，未并入超限块（内存上界 ≈ max + 单块）。"""
    from app.tools.parse_material import _collect_capped

    out, too_large = asyncio.run(_collect_capped(_FakeAiter([b"x" * 8, b"x" * 8, b"y"]), 10))
    assert too_large is True
    assert out == b"x" * 8  # 第二块会令 total=16>10，未并入


def test_collect_capped_exact_boundary_ok():
    """恰好等于上限不视为超限。"""
    from app.tools.parse_material import _collect_capped

    out, too_large = asyncio.run(_collect_capped(_FakeAiter([b"x" * 10]), 10))
    assert out == b"x" * 10
    assert too_large is False


def test_decode_best_utf8_and_no_crash():
    """utf-8 精确解码；非 utf-8 字节不崩（返回 str，具体编码依赖 charset_normalizer 检测）。"""
    from app.tools.parse_material import _decode_best

    assert _decode_best(b"") == ""
    assert _decode_best("你好".encode("utf-8")) == "你好"
    # GBK 字节不应崩（charset_normalizer 检测；短串可能不稳，仅断言返回 str）
    assert isinstance(_decode_best("你好".encode("gbk")), str)
    # 纯 latin-1 字节兜底不崩
    assert isinstance(_decode_best(b"\xff\xfe\x80"), str)


# --- zip bomb 防护（R23，plan/09 §四，iter 76）---


def test_zip_safety_rejects_huge_uncompressed(tmp_path, monkeypatch):
    import zipfile

    import app.tools.parse_material as pm

    monkeypatch.setattr(pm, "_ZIP_MAX_UNCOMPRESSED", 1024)  # 调低上限避免真写 512MB
    p = tmp_path / "bomb.docx"
    with zipfile.ZipFile(p, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("big.xml", b"x" * 2048)  # 解压 2KB > 上限 1KB

    with pytest.raises(pm._ParseError, match="超上限"):
        pm._check_zip_safety(p)


def test_zip_safety_rejects_too_many_entries(tmp_path):
    import zipfile

    from app.tools.parse_material import _ZIP_MAX_ENTRIES, _check_zip_safety, _ParseError

    p = tmp_path / "many.pptx"
    with zipfile.ZipFile(p, "w") as zf:
        for i in range(_ZIP_MAX_ENTRIES + 1):
            zf.writestr(f"f{i}", "x")

    with pytest.raises(_ParseError, match="条目数"):
        _check_zip_safety(p)


def test_zip_safety_bad_zip(tmp_path):
    from app.tools.parse_material import _check_zip_safety, _ParseError

    p = tmp_path / "broken.docx"
    p.write_bytes(b"not a zip at all")

    with pytest.raises(_ParseError, match="损坏"):
        _check_zip_safety(p)


def test_zip_safety_normal_docx_passes(tmp_path):
    import zipfile

    from app.tools.parse_material import _check_zip_safety

    p = tmp_path / "ok.docx"
    with zipfile.ZipFile(p, "w") as zf:
        zf.writestr("word/document.xml", "<xml/>")
    _check_zip_safety(p)  # 不抛即通过


def test_url_unexpected_error_not_returned_as_content(monkeypatch):
    """URL 抓取抛非预期异常时降级为空 -> empty_content，而非把 URL 串当解析文本假成功。"""
    import app.tools.parse_material as pm

    async def _boom(source):
        raise OSError("network stack exploded")

    monkeypatch.setattr(pm, "_parse_url", _boom)
    out = _run({"kind": "url", "source": "https://example.com/a"})
    assert out.status == "error"
    assert out.errors[0].code == "empty_content"
